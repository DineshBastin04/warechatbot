"""Test fakes and builders. Nothing here touches a real database, mail server,
LLM or external API."""
import functools
import json
import os
import sys
import tempfile
from datetime import datetime
from zoneinfo import ZoneInfo

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from engine.warehouse_perf.delivery import EmailDelivery  # noqa: E402
from engine.warehouse_perf.service import ServiceSettings, WarehousePerfService  # noqa: E402
from engine.warehouse_perf.store import WarehousePerfStore  # noqa: E402
from engine.warehouse_perf.synthetic import (FakeWorkspaceCollection, SyntheticData, SyntheticDatabase,  # noqa: E402
                                             generate_tables)
from engine.warehouse_perf.workspace_context import WorkspaceResolver  # noqa: E402

TZ = ZoneInfo("Asia/Kolkata")
NOW = datetime(2026, 9, 24, 13, 30)          # a Thursday, naive local time
WS_A = "aaaaaaaa-0000-4000-8000-00000000000a"
WS_B = "bbbbbbbb-0000-4000-8000-00000000000b"
DB_PASSWORD = "db-secret-pw-123"

_TABLE_CACHE = {}


def tables(**kw):
    """Cached synthetic tables (generation is the slow part)."""
    key = json.dumps({k: (str(v) if isinstance(v, datetime) else v) for k, v in sorted(kw.items())}, default=str)
    if key not in _TABLE_CACHE:
        args = {"now": NOW, **kw}
        _TABLE_CACHE[key] = generate_tables(**args)
    return {k: v.copy() for k, v in _TABLE_CACHE[key].items()}


def ws_meta(name, server, recipients=None, llm=None, db="WMS_DB"):
    return {
        "name": name, "display_name": name,
        "db_config": json.dumps({"serverName": server, "port": "1433", "databaseName": db, "username": "u_" + name,
                                 "password": DB_PASSWORD}),
        "db_config_b": json.dumps({}),
        "llm_config": json.dumps(llm or {}),
        "email_recipients": json.dumps(recipients if recipients is not None else
                                       [{"label": "Operations Head", "address": f"ops@{name.lower()}.example"}]),
    }


class MutableCollection(FakeWorkspaceCollection):
    def update_meta(self, ws, **fields):
        self._ws[ws] = {**self._ws[ws], **fields}

    def add(self, ws, meta):
        self._ws[ws] = meta

    def delete(self, ws):
        self._ws.pop(ws, None)


class BrokenCollection:
    def get(self, ids=None):
        raise RuntimeError("chroma unavailable")


def two_workspaces():
    return MutableCollection([
        {"id": WS_A, "metadata": ws_meta("Alpha", "srv-alpha")},
        {"id": WS_B, "metadata": ws_meta("Beta", "srv-beta",
                                         recipients=[{"label": "Operations Head", "address": "ops@beta.example"},
                                                     {"label": "Beta Only", "address": "only@beta.example"}])},
    ])


class RecordingSender:
    def __init__(self, fail_times=0, result=True):
        self.calls = []
        self.fail_times = fail_times
        self.result = result

    def __call__(self, recipients, subject, body, rows, filename, html=None, inline_images=None,
                 extra_attachments=None):
        self.calls.append({"recipients": list(recipients), "subject": subject, "body": body, "rows": rows,
                           "filename": filename, "html": html, "inline_images": inline_images or [],
                           "extra_attachments": extra_attachments or []})
        if self.fail_times > 0:
            self.fail_times -= 1
            raise ConnectionError(f"SMTP login failed for password={DB_PASSWORD}")
        return self.result


class EmailConfig:
    def __init__(self, enabled=True):
        self.enabled = enabled

    def __call__(self):
        return {"send_email_enabled": self.enabled, "email_provider": "gmail", "sender_email": "bot@example.com",
                "sender_password": "smtp-secret"}


class DataByServer:
    """data_factory: maps each workspace's configured server to its own synthetic DB,
    and records which workspace asked for data."""

    def __init__(self, dbs, **data_kw):
        self.dbs = dbs
        self.data_kw = data_kw
        self.requests = []
        self.instances = []

    def __call__(self, ctx):
        server = ctx.db_config.get("serverName")
        self.requests.append((ctx.workspace_id, server))
        kw = self.data_kw.get(server, {})
        data = SyntheticData(self.dbs[server], **kw)
        self.instances.append((ctx.workspace_id, data))
        return data


def make_service(coll=None, dbs=None, sender=None, email_cfg=None, now=NOW, narrator=None, data_kw=None,
                 store_dir=None, **settings):
    coll = coll or two_workspaces()
    dbs = dbs or {"srv-alpha": SyntheticDatabase(tables()), "srv-beta": SyntheticDatabase(tables(degrade=False,
                                                                                                 seed=11))}
    store = WarehousePerfStore(store_dir or tempfile.mkdtemp(prefix="wpa-test-"))
    sender = sender or RecordingSender()
    factory = DataByServer(dbs, **(data_kw or {}))
    svc = WarehousePerfService(
        store=store, resolver=WorkspaceResolver(coll), executor=None,
        delivery=EmailDelivery(store, sender, email_cfg or EmailConfig(), sleep=lambda s: None,
                               label_filter=settings.pop("label_filter", None)),
        settings=ServiceSettings(timezone="Asia/Kolkata", **settings),
        clock=(lambda: now.replace(tzinfo=TZ)) if not callable(now) else now,
        narrator_factory=(lambda cfg: narrator) if narrator is not None else (lambda cfg: None),
        data_factory=factory)
    return svc, sender, store, factory, coll


class FakeNarrator:
    def __init__(self, text=None, exc=None):
        self.text = text
        self.exc = exc
        self.prompts = []

    def complete(self, system_prompt, user_prompt):
        self.prompts.append(user_prompt)
        if self.exc:
            raise self.exc
        return self.text


# ------------------------------------------------------------ flask app
def make_flask(svc):
    """A bare Flask app with stand-in auth decorators: the test user is
    passed as JSON in the X-Test-User header."""
    from flask import Flask, jsonify, request

    from engine.warehouse_perf.chat import WarehouseChat
    from engine.warehouse_perf.routes import register_routes

    app = Flask("wpa-test")

    def requires_auth(f):
        @functools.wraps(f)
        def wrapper(*a, **kw):
            raw = request.headers.get("X-Test-User")
            if not raw:
                return "login", 302
            kw["user"] = json.loads(raw)
            return f(*a, **kw)
        return wrapper

    def requires_role(roles):
        def deco(f):
            @functools.wraps(f)
            def wrapper(*a, **kw):
                if (kw.get("user") or {}).get("role") not in roles:
                    return jsonify({"type": "error", "error": "forbidden"}), 403
                return f(*a, **kw)
            return wrapper
        return deco

    register_routes(app, requires_auth, requires_role, svc, WarehouseChat(svc))
    return app


def user(role="admin", workspace="", username=None):
    return json.dumps({"username": username or f"{role}-user", "role": role, "workspace": workspace})


def docx_text(data):
    """All paragraph + table text of a .docx (bytes)."""
    import io
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return chr(10).join(parts)
