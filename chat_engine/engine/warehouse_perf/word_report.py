"""
Word (.docx) version of a report — the file attached to the email and offered
by the dashboard's "Download Word" button.

Same content as the Full report page, built only from the structured result:
  1. management summary: status + headline, "In short", KPI table, why it
     happened (the Advisor's likely causes with evidence), what to do now, the
     three charts (PNG, from charts.py), what was ruled out, what we don't know yet;
  2. technical details: the report sections (reporting.sections).
"""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import exec_report as X

WORD_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
INK = RGBColor(0x1A, 0x1A, 0x19)
INK_3 = RGBColor(0x76, 0x75, 0x6F)
STATE_WORD = {"worse": "Worse than normal", "better": "Better than normal", "normal": "Normal"}


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _left_border(cell, hex_colour, size=36):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:color"), hex_colour.lstrip("#"))
    borders.append(left)
    tc_pr.append(borders)


def _run(par, text, bold=False, size=None, colour=None, italic=False):
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if colour is not None:
        r.font.color.rgb = colour
    return r


def _bullets(doc, items):
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def build_docx(report):
    r = report["result"]
    ov = r.get("overall") or {}
    cls = ov.get("classification") if r.get("status") == "ok" else "insufficient_data"
    status, colour, icon = X.STATUS.get(cls, X.STATUS["insufficient_data"])
    names = X._plain(r)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    t = doc.add_paragraph()
    _run(t, "Warehouse Performance Report", bold=True, size=20, colour=INK)
    sub = doc.add_paragraph()
    _run(sub, f"{report.get('workspace_name') or ''} · {(r.get('period') or {}).get('label') or ''} · "
              f"report {report.get('report_id')}", size=9.5, colour=INK_3)

    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = banner.rows[0].cells[0]
    _shade(cell, "F7F7F5")
    _left_border(cell, colour)
    p = cell.paragraphs[0]
    _run(p, f"{icon} {status}", bold=True, size=10.5, colour=INK)
    p2 = cell.add_paragraph()
    _run(p2, X.headline(r), bold=True, size=14, colour=INK)

    ai = doc.add_table(rows=1, cols=1)                  # AI notice, right under the headline
    ai.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = ai.rows[0].cells[0]
    _shade(cell, "FFF8E6")
    _left_border(cell, "E0A82E")
    p = cell.paragraphs[0]
    _run(p, "AI-generated report: ", bold=True, size=9.5, colour=INK)
    _run(p, X.AI_NOTICE.split(": ", 1)[1], size=9.5, colour=INK)

    doc.add_heading("In short", level=2)
    _bullets(doc, X.summary_points(r))

    kpis = X._kpis(r)
    if kpis:
        doc.add_heading("Key numbers", level=2)
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        for i, h in enumerate(("Measure", "Today", "Normal", "Change", "Status")):
            hc = tbl.rows[0].cells[i]
            _shade(hc, "E8EEF8")
            hc.paragraphs[0].text = ""
            _run(hc.paragraphs[0], h, bold=True)
        for m in kpis:
            row = tbl.add_row().cells
            pct = m.get("pct_diff")
            arrow = "▼ " if (pct or 0) < 0 else "▲ " if (pct or 0) > 0 else ""
            vals = (names.get(m["metric_id"], m["label"]), X._num(m.get("actual"), m.get("unit")),
                    X._num(m.get("baseline"), m.get("unit")), arrow + X._pct(pct), STATE_WORD[X._state(m)])
            for i, v in enumerate(vals):
                row[i].text = str(v)

    diag = X.diagnosis(r) if r.get("status") == "ok" else None
    if diag:
        doc.add_heading("Why it happened", level=2)
        if diag.get("likely_causes"):
            if diag["status"] != "cause_found":
                doc.add_paragraph(diag["summary"])
            for c in diag["likely_causes"]:
                box = doc.add_table(rows=1, cols=1)
                cell = box.rows[0].cells[0]
                _shade(cell, "FAFAF8")
                _left_border(cell, X.CONF_COLOUR[c["confidence"]], size=28)
                p = cell.paragraphs[0]
                _run(p, f"#{c['rank']} · {X.CONF_WORD[c['confidence']]} · about {c['explains_pct']}% of the drop",
                     size=8.5, colour=INK_3)
                _run(cell.add_paragraph(), c["title"], bold=True, size=11.5, colour=INK)
                for e in c["evidence"][:4]:
                    _run(cell.add_paragraph(), "• " + e, size=9.5, colour=INK)
                if c.get("note"):
                    _run(cell.add_paragraph(), c["note"], size=9, colour=INK_3, italic=True)
                doc.add_paragraph()
        else:
            doc.add_paragraph(diag.get("summary") or "")

    recs = X.recommendations(r) if r.get("status") == "ok" else []
    if recs:
        doc.add_heading("What to do now", level=2)
        _bullets(doc, recs)

    for _cid, png, title, ans, note, _alt in (X._chart_specs(r) if r.get("status") == "ok" else []):
        doc.add_heading(title, level=2).paragraph_format.keep_with_next = True
        a = doc.add_paragraph()                             # the answer first, in words
        a.paragraph_format.keep_with_next = True
        _run(a, ans, bold=True, size=12, colour=INK)
        cap = doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True          # heading + text stay on the chart's page
        _run(cap, note, size=9.5, colour=INK_3)
        doc.add_picture(io.BytesIO(png), width=Cm(16.5))

    if (diag or {}).get("ruled_out"):
        doc.add_heading("Checked and ruled out", level=2)
        _bullets(doc, [x["statement"] for x in diag["ruled_out"][:8]])
    unknown = X.unknown_text(r)
    if unknown:
        doc.add_heading("What we don't know yet", level=2)
        _bullets(doc, [unknown] + [w["message"] for w in r.get("freshness_warnings") or []])

    note = doc.add_paragraph()
    _run(note, "“Normal” means the same weekday over the previous five weeks, over the same hours. A change is "
               "flagged only when it is well outside the usual day-to-day variation. Three agents produced this "
               "report: the Monitor (what changed), the Investigator (checked the possible causes in the warehouse "
               "records) and the Advisor (ranked them and wrote the actions). They only read the warehouse database "
               "and change nothing.", size=8.5, colour=INK_3, italic=True)
    _run(doc.add_paragraph(), X.AI_DISCLAIMER, size=8.5, colour=INK, bold=True)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Technical details", level=1)
    from .reporting import sections   # local import: reporting imports this module's siblings
    for title, lines in sections(report):
        doc.add_heading(title, level=2)
        for line in lines:
            if not line:
                continue
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
